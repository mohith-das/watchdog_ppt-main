import numpy as np
import pandas as pd
from bigquery import get_bigquery_client
from config import REVERSE_METRICS
from pptx.util import Inches, Cm, Pt
from pptx.dml.color import RGBColor


def delta_pct(now, prev):
    if np.isnan(prev):
        prev = 0

    if np.isnan(now):
        now = 0

    if prev == 0:
        if now > 0:
            return float('inf')
        elif now < 0:
            return -float('inf')
        else:
            return 0
    else:
        return (now / prev - 1) * 100


def delta_print_pct(prev, now):
    prev = float(prev)
    now = float(now)
    if prev == 0:
        if now == 0:
            return 0
        else:
            return "∞"
    else:
        result = now / prev - 1
        if abs(result) < 10:
            return abs(round((now / prev - 1) * 100, 2))
        else:
            return ">1000"


def print_delta(prev, now):
    greater = now > prev
    return f"{'▲' if greater else '▼'}{delta_print_pct(prev=prev, now=now)}%"
    # return f"{'+' if greater else '-'}{delta_print_pct(prev=prev, now=now)}%"


def human_format(num):
    if num == np.inf:
        return '∞'

    if num == -np.inf:
        return '-∞'

    sign = '-' if num < 0 else ''

    num = abs(num)
    magnitude = 0
    original_num = num
    while abs(num) >= 1000:
        magnitude += 1
        num /= 1000.0
    # add more suffixes if you need them
    if (magnitude == 1):
        return f"{sign}{(round(num, 2))}K"
    elif (magnitude == 2):
        return f"{sign}{round(num, 2)}M"
    else:
        return f"{sign}{int(original_num)}"


def fix_name(metric):
    if metric.lower() == 'ctr':
        return 'Click Through Rate'

    if metric.lower() == 'cpc':
        return 'Cost Per Click'

    if '__to__' in metric:
        metric = ' '.join(metric.split('_'))
        result = []
        for word in metric.split():
            result.append(word.title())
        return ' '.join(result) + ' CVR'
    else:
        return ' '.join(metric.split('_'))


def get_metric_type(metric):
    metric = metric.lower()

    if metric.startswith('revenue') or metric.endswith('revenue') or metric.endswith('spend') or metric.endswith(
            'cost'):
        return 'revenue'
    elif metric.startswith('aov') or metric.startswith('cpc') or metric.startswith('cac'):
        return 'aov'
    elif metric.endswith('rate') or ('__to__' in metric) or metric == 'ctr' or metric.endswith(
            'share') or metric.lower() == 'acos':
        return 'rate'
    else:
        return 'traffic'



def print_formatted(value, metric):
    metric = metric.lower()

    if np.isnan(value):
        return None

    metric_type = get_metric_type(metric)

    if metric_type == 'revenue':
        return f"${human_format(value)}"
    elif metric_type == 'aov':
        return f"${value:.2f}"
    elif metric_type == 'rate':
        return f"{100*value:.2f}%"
    else:
        return f"{human_format(value)}"


def divide(a, b):
    if b != 0:
        return a/b
    else:
        return 0


def filter_dims(kpi_df, dim_values_to_include):
    for dim in dim_values_to_include:
        mask = kpi_df[dim].isin(dim_values_to_include[dim])
        kpi_df = kpi_df[mask]

    return kpi_df


def not_none(number):
    return (number == number) and number is not None


def is_none(number):
    return not not_none(number)


def get_tables(project_id, dataset_id, period):
    result = []
    client = get_bigquery_client(project_id)
    table_ids = [table.table_id for table in client.list_tables(dataset_id)]
    for table_id in table_ids:
        if table_id.endswith('_view') and (period in table_id) and ('raw_funnel' not in table_id):
            result.append(table_id[:-5] + '_anomaly')

    return result


def get_dim_metrics(project_id, dataset_id, table_id):
    client = get_bigquery_client(project_id)
    dataset_ref = client.dataset(dataset_id, project=project_id)
    table_ref = dataset_ref.table(table_id)
    table = client.get_table(table_ref)

    dims = []
    metrics = []
    for schema in table.schema:
        if schema.field_type in ['NUMERIC', 'FLOAT', 'INTEGER']:
            if schema.name.endswith('_yhat'):
                metrics.append(schema.name[:-5])
        elif schema.name.lower() in ['date', 'datehour', 'week']:
            continue
        else:
            dims.append(schema.name)

    if len(dims) == 0:
        return None, metrics

    elif len(dims) == 1:
        return dims[0], metrics

    else:
        raise ValueError(f'Error while getting dim and metrics for {project_id}.{dataset_id}.{table_id} : length of dims is {len(dims)}')


def get_anomaly_df(project_id, dataset_id, table_id, period, date_filter=None):
    if period == 'hourly':
        date_col = 'DateHour'
    elif period == 'daily':
        date_col = 'Date'
    elif period == 'weekly':
        date_col = 'Week'
    else:
        raise Exception(f"Invalid period - {period}")

    if date_filter:
        query = f"""
                SELECT * FROM
                `{project_id}.{dataset_id}.{table_id}`
                WHERE {date_col} {date_filter}
                ORDER BY {date_col}
                """
    else:
        query = f"""
                SELECT * FROM
                `{project_id}.{dataset_id}.{table_id}`
                ORDER BY {date_col}
                """

    client = get_bigquery_client(project_id)
    anomaly_df = (
        client.query(query)
            .result()
            .to_dataframe()
    )

    _, metrics = get_dim_metrics(project_id, dataset_id, table_id)

    for metric in metrics:
        anomaly_df[metric] = pd.to_numeric(anomaly_df[metric])

    return anomaly_df


def get_table_details(project_id, asset, table_id):
    data_source = get_data_source(table_id)
    dim, metrics = get_dim_metrics(project_id, asset, table_id)

    return data_source, dim, metrics


def get_anomaly_type(y, upper, lower):
    if y is None:
        return 0

    if y > upper:
        return 1
    elif y < lower:
        return -1
    else:
        return 0


def check_warning(y, upper, lower, lower_bound=10, upper_bound=30):
    return (lower_bound < -delta_pct(y, lower) <= upper_bound) or (lower_bound < delta_pct(y, upper) <= upper_bound)


def check_critical(y, upper, lower, threshold=30):
    return (-delta_pct(y, lower) > threshold) or (delta_pct(y, upper) > threshold)


def get_color(anomaly_type, metric):
    if metric.lower() in REVERSE_METRICS:
        anomaly_type *= -1

    if anomaly_type == 1:
        return 'green'
    elif anomaly_type == -1:
        return 'red'
    else:
        return None


def get_data_source(table_id):
    if table_id.startswith('ga_'):
        data_source = 'Google Analytics'
    elif table_id.startswith('fb_'):
        data_source = 'Facebook'
    elif table_id.startswith('googleAds_'):
        data_source = 'Google Ads'
    elif table_id.startswith('custom_') or table_id.startswith('shopify_') or table_id.startswith('ecommerce_'):
            data_source = 'Ecommerce'
    elif table_id.startswith('magento_'):
            data_source = 'Magento'
    elif table_id.startswith('upscribe_'):
        data_source = 'Upscribe'
    elif table_id.startswith('affiliate_'):
        data_source = 'Affiliate'
    else:
        data_source = table_id.split('_')[0]

    return data_source


def add_color(txt, color):
    if color:
        # return f"""<span style="color:{color}">{txt}</span>"""
        return f"""<font color="{color}">{txt}</font>"""
    else:
        return txt


class Element:
    def __init__(self, text, type='regular_text'):
        self.text = text
        self.type = type


def bold(txt):
    return f'<b>{txt}</b>'


def main_heading(text):
    text = bold(text)
    # return f'<h1 style="font-size:25px">{text}</h1><br>'
    # return f"<h1>{add_color(text, '#2980b9')}</h1><br>"
    return f"<h1>{add_color(text, '#2980b9')}</h1>"


def sub_heading(text):
    text = bold(text)
    # return f'<br><h1 style="font-size:23px">{text}</h1><br>'
    # return f"<h3>{add_color(text, '#2980b9')}</h3><br>"
    return f"<h3>{add_color(text, '#2980b9')}</h3>"


def regular_text(text):
    return f'<p>{text}</p>'


always_include_data_source = ['Google Ads', 'Facebook']


def print_anomaly(p, row, pre):

    data_source = row.data_source
    dimension = row.dimension
    dim_label = row.dim_label
    metric = row.metric
    period = row.period
    is_warning = row.is_warning
    is_critical = row.is_critical
    y = row.y
    y_prev = row.y_prev
    yhat = row.yhat
    yhat_upper = row.yhat_upper
    yhat_lower = row.yhat_lower

    fact_vs_forecast = ''

    prev_anomaly_type = row.anomaly_type
    prev_color = get_color(prev_anomaly_type, metric)

    yhat_anomaly_type = get_anomaly_type(y=y, upper=yhat_upper, lower=yhat_lower)
    yhat_color = get_color(yhat_anomaly_type, metric)

    if data_source in always_include_data_source:
        fact_vs_forecast = fact_vs_forecast + f'{data_source} '

    if not_none(dim_label):
        # fact_vs_forecast = fact_vs_forecast + bold(f'"{fix_name(dimension)}: {fix_name(dim_label)}"')
        fact_vs_forecast = fact_vs_forecast + f'"{fix_name(dim_label)}"'
    fact_vs_forecast = fact_vs_forecast + f" {fix_name(metric)} {print_formatted(y, metric)}"

    run = p.add_run()
    run.text = '\n' + f'{pre}{fact_vs_forecast}'

    font = run.font
    font.name = 'Poppins'
    font.size = Pt(12)
    font.bold = True
    font.color.rgb = RGBColor(100, 100, 100)

    if yhat_color == 'red':
        run = p.add_run()
        run.text = f" ({print_delta(now=y, prev=yhat)})"

        font = run.font
        font.name = 'Poppins'
        font.size = Pt(12)
        font.bold = True
        font.color.rgb = RGBColor(200, 0, 0)
    elif yhat_color == 'green':
        run = p.add_run()
        run.text = f" ({print_delta(now=y, prev=yhat)})"

        font = run.font
        font.name = 'Poppins'
        font.size = Pt(12)
        font.bold = True
        font.color.rgb = RGBColor(0, 200, 0)
    else:
        run = p.add_run()
        run.text = f" ({print_delta(now=y, prev=yhat)})"

        font = run.font
        font.name = 'Poppins'
        font.size = Pt(12)
        font.bold = True
        font.color.rgb = RGBColor(100, 100, 100)


def print_comment(row):
    dimension = row.dimension
    dim_label = row.dim_label
    metric = row.metric
    period = row.period
    is_warning = row.is_warning
    is_critical = row.is_critical
    y = row.y
    y_prev = row.y_prev
    yhat = row.yhat
    yhat_upper = row.yhat_upper
    yhat_lower = row.yhat_lower

    if period == 'daily':
        comparison = 'SDLW'
    elif period == 'weekly':
        comparison = 'WoW'
    else:
        raise Exception(f"Invalid period - {period}")

    comment = ''

    if not_none(dim_label):
        # fact_vs_forecast = fact_vs_forecast + bold(f'"{fix_name(dimension)}: {fix_name(dim_label)}"')
        comment = comment + f'"{fix_name(dim_label)}"'
    comment = comment + f" {fix_name(metric)}"

    if y > y_prev:
        comment = comment + f" increased by "
    else:
        comment = comment + f" decreased by "

    comment = comment + f"{print_formatted(y-y_prev, metric)}"

    comment = comment + f" {comparison}"

    return comment


def filter_data_by_kpi(asset_df, kpi):
    data_source_mask = asset_df['data_source'] == kpi.data_source

    if kpi.dimension:
        dimension_mask = asset_df['dimension'] == kpi.dimension
    else:
        dimension_mask = asset_df['dimension'].isnull()

    if kpi.dim_label:
        dim_label_mask = asset_df['dim_label'] == kpi.dim_label
    else:
        dim_label_mask = asset_df['dim_label'].isnull()

    metric_mask = asset_df['metric'] == kpi.metric

    filtered_df = asset_df[data_source_mask & dimension_mask & dim_label_mask & metric_mask]
    return filtered_df

